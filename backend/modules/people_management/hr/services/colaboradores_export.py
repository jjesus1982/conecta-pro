"""Exportar / imprimir colaboradores e contar por status — DGX V4 (24/09/2026), lacuna #22.

A DGX tem `ListagemReport` (PDF/Excel/Word) em TODA tela e, no `/Colaboradores/Index`, uma
faixa de contadores por status. Aqui não havia NENHUM export de colaboradores: para mandar a
lista ao contador, alguém copiava a tela e colava no Excel.

Réguas IMPORTADAS, nunca recriadas — é o ponto todo desta frente:
  · `identidade.SEM_VINCULO`  → quem não é gente da casa (inativo/candidato/pj_pendente);
  · `mapa_de_ponto.COORTE`    → quem entra na conta do ponto (ativo e não homologação);
  · `coorte_ponto.SQL_NAO_AUSENTE_HOJE` → férias aprovadas de hoje + afastamento aberto +
    desligamento na reta final. «Ausente hoje» é a negação dela dentro da coorte.
Duas cópias da mesma régua divergem, e a que diverge cala (lição de 11/09, `identidade.py`).

LGPD: o CPF sai MASCARADO no PDF sempre (documento circula em papel e por e-mail) e completo
no Excel só para admin — `admin=False` mascara lá também. O PDF reusa `gerar_tabela_pdf`
(paisagem, cabeçalho repetido por página, timbrado `pdf_branding`), o mesmo do relatório de
diaristas: um relatório de lista se confere linha a linha, e isso é tabela.
"""

from __future__ import annotations

import csv
import io
from datetime import date

from sqlalchemy import text

from modules.integrations.connectors.whatsapp.identidade import SEM_VINCULO
from modules.people_management.ponto.coorte_ponto import SQL_NAO_AUSENTE_HOJE
from modules.people_management.ponto.mapa_de_ponto import COORTE

#: Vínculo vivo: nem sem-vínculo, nem demitido, nem cadastro de homologação.
COM_VINCULO = (
    "lower(coalesce(e.status,'')) <> ALL(:sem) AND lower(coalesce(e.status,'')) <> 'demitido' "
    "AND coalesce(e.is_homologacao,false) = false"
)
_AFASTADO = (
    "EXISTS (SELECT 1 FROM sst_afastamentos a WHERE a.employee_id = e.id "
    "AND lower(coalesce(a.status,'')) IN ('em_andamento','ativo') AND a.data_retorno IS NULL)"
)
_FERIAS = (
    "EXISTS (SELECT 1 FROM hr_vacation_requests v WHERE v.employee_id = e.id "
    "AND upper(coalesce(v.status,'')) = 'APPROVED' "
    "AND (now() AT TIME ZONE 'America/Manaus')::date BETWEEN v.start_date AND v.end_date)"
)

#: (rótulo, predicado SQL, é subconjunto de «Ativo»?). A tela e o oráculo leem DAQUI — um
#: contador cuja régua só existe dentro do SELECT é um número que ninguém consegue recontar.
CONTADORES: tuple[tuple[str, str, bool], ...] = (
    ("Ativo", COM_VINCULO, False),
    ("Inativo", "lower(coalesce(e.status,'')) = ANY(:sem)", False),
    ("Demitido", "lower(coalesce(e.status,'')) = 'demitido'", False),
    ("Suspenso", "lower(coalesce(e.status,'')) = 'suspenso'", True),
    ("Afastado", f"{COM_VINCULO} AND {_AFASTADO}", True),
    ("Férias", f"{COM_VINCULO} AND {_FERIAS}", True),
    ("Ausente hoje", f"{COORTE} AND NOT (TRUE {SQL_NAO_AUSENTE_HOJE})", True),
)

#: Colunas do Excel/CSV — o núcleo da ficha da T1 (dados gerais · contrato · documentos ·
#: endereço/contato · banco), achatado em uma linha por pessoa.
COLUNAS: tuple[tuple[str, str], ...] = (
    ("Matrícula", "matricula"),
    ("Nome", "nome"),
    ("CPF", "cpf"),
    ("Nascimento", "data_nascimento"),
    ("Sexo", "sexo"),
    ("Estado civil", "estado_civil"),
    ("Nome da mãe", "nome_mae"),
    ("Status", "status"),
    ("Cargo", "cargo"),
    ("Função (CCT)", "funcao_cct"),
    ("CBO", "cbo"),
    ("Departamento", "departamento"),
    ("Centro de custo", "centro_custo"),
    ("Admissão", "data_admissao"),
    ("Demissão", "data_demissao"),
    ("Tipo de contrato", "tipo_contrato"),
    ("Escala", "escala_padrao"),
    ("Jornada", "jornada_trabalho"),
    ("Salário base", "salario_base"),
    ("Insalubridade %", "insalubridade_percentual"),
    ("Periculosidade %", "periculosidade_percentual"),
    ("Posto atual", "posto_atual_nome"),
    ("Cliente/Condomínio", "condominio_nome"),
    ("PIS", "pis"),
    ("CTPS", "ctps_numero"),
    ("RG", "rg"),
    ("CNH", "cnh_numero"),
    ("CNV", "cnv"),
    ("E-mail", "email"),
    ("Celular", "celular"),
    ("Cidade/UF", "cidade_uf"),
    ("CEP", "cep"),
    ("Banco", "banco"),
    ("Agência", "agencia"),
    ("Conta", "conta"),
    ("Chave PIX", "pix"),
)

_SQL_LINHAS = """
SELECT
  coalesce(e.matricula,'') AS matricula, e.nome, coalesce(e.cpf,'') AS cpf, e.data_nascimento,
  coalesce(e.sexo,'') AS sexo, coalesce(e.estado_civil,'') AS estado_civil, coalesce(e.nome_mae,'') AS nome_mae,
  coalesce(e.status,'') AS status, coalesce(e.cargo,'') AS cargo,
  -- `cbo` nasce no `_ensure` da T1 (ALTER TABLE ... ADD COLUMN IF NOT EXISTS). Lido por
  -- to_jsonb para a exportação não depender da ORDEM em que os `_ensure` rodaram.
  coalesce(cc.cargo_nome,'') AS funcao_cct, coalesce(to_jsonb(cc) ->> 'cbo','') AS cbo,
  coalesce(e.departamento,'') AS departamento, coalesce(e.centro_custo,'') AS centro_custo,
  e.data_admissao, e.data_demissao, coalesce(e.tipo_contrato,'') AS tipo_contrato,
  coalesce(e.escala_padrao,'') AS escala_padrao, coalesce(e.jornada_trabalho,'') AS jornada_trabalho,
  e.salario_base, e.insalubridade_percentual, e.periculosidade_percentual,
  coalesce(e.posto_atual_nome,'') AS posto_atual_nome, coalesce(cd.nome,'') AS condominio_nome,
  coalesce(e.pis,'') AS pis, coalesce(e.ctps_numero,'') AS ctps_numero, coalesce(e.rg,'') AS rg,
  coalesce(e.cnh_numero,'') AS cnh_numero, coalesce(e.cnv,'') AS cnv,
  coalesce(e.email,'') AS email, coalesce(e.celular, e.telefone,'') AS celular,
  trim(both '/' from concat_ws('/', e.cidade, e.uf)) AS cidade_uf, coalesce(e.cep,'') AS cep,
  coalesce(e.banco,'') AS banco, coalesce(e.agencia,'') AS agencia, coalesce(e.conta,'') AS conta,
  coalesce(e.pix_key, e.pix, '') AS pix
FROM employees e
LEFT JOIN cct_cargos cc ON cc.id = e.cct_cargo_id
LEFT JOIN condominios cd ON cd.id = e.cliente_id
WHERE {onde}
ORDER BY e.nome
"""


def mascara_cpf(cpf: str | None) -> str:
    """000.***.***-00 — o bastante para conferir, insuficiente para vazar."""
    d = "".join(c for c in str(cpf or "") if c.isdigit())
    if len(d) != 11:
        return str(cpf or "")
    return f"{d[:3]}.***.***-{d[9:]}"


def _fmt(v) -> str:
    if v is None:
        return ""
    if isinstance(v, date):
        return v.strftime("%d/%m/%Y")
    return str(v)


def _onde(status: str | None, condominio: str | None, funcao: str | None) -> tuple[str, dict]:
    """Filtro da lista. `status` vazio = só quem tem vínculo vivo (o default da tela do DP);
    `status='todos'` = o quadro inteiro, inclusive demitido — é o que o contador pede."""
    params: dict = {"sem": list(SEM_VINCULO)}
    s = (status or "").strip().lower()
    if not s or s == "ativo":
        onde = [COM_VINCULO]
    elif s == "todos":
        onde = ["TRUE"]
    elif s == "afastado":
        onde = [COM_VINCULO, _AFASTADO]
    elif s == "ferias":
        onde = [COM_VINCULO, _FERIAS]
    elif s == "inativo":
        onde = ["lower(coalesce(e.status,'')) = ANY(:sem)"]
    else:
        onde = ["lower(coalesce(e.status,'')) = :st"]
        params["st"] = s
    if (condominio or "").strip():
        onde.append("(cd.nome ILIKE :cond OR e.posto_atual_nome ILIKE :cond)")
        params["cond"] = f"%{condominio.strip()}%"
    if (funcao or "").strip():
        onde.append("(e.cargo ILIKE :fn OR cc.cargo_nome ILIKE :fn)")
        params["fn"] = f"%{funcao.strip()}%"
    return " AND ".join(f"({o})" for o in onde), params


async def contadores(db) -> list[dict]:
    """Um SELECT, um FILTER por contador — recontável linha a linha por quem duvidar."""
    sel = ", ".join(f"count(*) FILTER (WHERE {pred})" for _, pred, _ in CONTADORES)
    r = (await db.execute(text(f"SELECT {sel} FROM employees e"), {"sem": list(SEM_VINCULO)})).first()  # noqa: S608
    return [
        {"label": rot, "valor": int(r[i] or 0), "subconjunto_de_ativo": sub}
        for i, (rot, _p, sub) in enumerate(CONTADORES)
    ]


async def linhas(db, status=None, condominio=None, funcao=None, admin: bool = False) -> list[dict]:
    onde, params = _onde(status, condominio, funcao)
    rows = (await db.execute(text(_SQL_LINHAS.format(onde=onde)), params)).mappings().all()
    out = []
    for r in rows:
        d = {campo: _fmt(r[campo]) for _rot, campo in COLUNAS}
        if not admin:
            d["cpf"] = mascara_cpf(d["cpf"])
        out.append(d)
    return out


def para_csv(dados: list[dict]) -> bytes:
    """`;` e latin-1: é o CSV que o Excel brasileiro abre com dois cliques (mesma escolha do
    export Domínio da folha). utf-8 puro faz o Excel BR mostrar «JOSÃ‰» na primeira coluna."""
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";", lineterminator="\r\n")
    w.writerow([rot for rot, _ in COLUNAS])
    for d in dados:
        w.writerow([d[campo] for _rot, campo in COLUNAS])
    return buf.getvalue().encode("latin-1", errors="replace")


def para_xlsx(dados: list[dict], titulo: str = "Colaboradores") -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = titulo[:31]
    ws.append([rot for rot, _ in COLUNAS])
    azul, branco = PatternFill("solid", fgColor="1E3A5F"), Font(bold=True, color="FFFFFF")
    for c in ws[1]:
        c.fill, c.font, c.alignment = azul, branco, Alignment(vertical="center")
    for d in dados:
        ws.append([d[campo] for _rot, campo in COLUNAS])
    ws.freeze_panes = "A2"
    for i, (rot, campo) in enumerate(COLUNAS, start=1):
        largura = max([len(rot)] + [len(d[campo]) for d in dados[:500]] or [len(rot)])
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = min(max(largura + 2, 10), 42)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


#: Colunas do PDF (paisagem): o que se confere no papel. CPF sempre mascarado.
_PDF_COLS = (
    ("Matrícula", "matricula", 20, "L"),
    ("Colaborador", "nome", 62, "L"),
    ("CPF", "cpf", 32, "L"),
    ("Cargo", "cargo", 48, "L"),
    ("Admissão", "data_admissao", 22, "C"),
    ("Posto / Cliente", "posto_atual_nome", 52, "L"),
    ("Status", "status", 22, "C"),
)


def para_pdf(dados: list[dict], subtitulo: str, resumo: list[dict] | None = None) -> bytes:
    from modules.financial.services.relatorio_diaristas_pdf import gerar_tabela_pdf

    return gerar_tabela_pdf(
        titulo="Relação de colaboradores",
        subtitulo=subtitulo,
        colunas=[{"t": t_, "w": w, "a": a} for t_, _c, w, a in _PDF_COLS],
        linhas=[[mascara_cpf(d[c]) if c == "cpf" else (d[c] or "—") for _t, c, _w, _a in _PDF_COLS] for d in dados],
        rodape=[(x["label"], str(x["valor"])) for x in (resumo or [])],
        nota="CPF mascarado (LGPD). O Excel desta mesma lista traz o CPF completo para administradores.",
    )
