"""DGX V5 — Fiscal/Financeiro: cadastro fiscal completo do serviço, formas de pagamento,
limite por condição, centros de custo em árvore e relatórios em PDF/Excel (24/09/2026).

O que foi cavado ANTES de construir (sandbox = cópia de produção de 23/09) — três coisas que o
brief mandava criar já existiam, e criar de novo teria feito tabela-irmã morta:

  · **Forma de pagamento já tem tabela**: `payment_methods` (modelo `financial/models/payment_method.py`)
    com `code, name, payment_type, bank_account_id, ativo` — exatamente o cadastro pedido — e
    `payable_accounts.payment_method_id`, `payable_payments.payment_method_id`,
    `suppliers.default_payment_method_id` já apontam para ela. Estava VAZIA (0 linhas) e sem tela:
    a lacuna era a PORTA, não a tabela. Por isso aqui NÃO nasce `fin_formas_pagamento` — nasce a
    tela sobre `payment_methods`, e o `payment_method_id` que já existe passa a ter o que apontar.
    `gera_boleto` do brief = `payment_type='boleto'` (não vira coluna nova).
  · **Centro de custo já tem pai**: `fin_cost_centers.parent_id` + `level` + `full_path` existem
    desde sempre (tabela vazia, 0 linhas). `pai_id` seria uma segunda coluna para a mesma coisa.
    Aqui entra a TELA da árvore (que mantém `level`/`full_path` coerentes) e o vínculo
    categoria-do-extrato → centro, que é o que falta para o orçado × realizado falar por centro.
  · **A emissão de NFS-e NÃO lê o cadastro fiscal**: o XML nacional (`core/nfse_nacional.py:316-322`)
    traz `<cNBS>120032900</cNBS>` e `<cTribMun>100</cTribMun>` FIXOS, e o cTribNac vem de
    `dps.servico.codigo_tributacao_nacional` (do chamador), nunca de `fin_codigos_servico`. Manaus
    (`core/nfse_manaus.py:286`) usa `nfse.servico.codigo_servico`, idem. Logo NBS/CST/IBS/CBS aqui
    são CADASTRO — ligar na emissão mudaria o XML mandado ao fisco, e isso é decisão do dono
    (relatório §7). O oráculo trava exatamente isso: o XML de uma nota de teste sai byte-idêntico
    antes e depois desta frente.

Prefixo `_` = o discovery pula; `financeiro.py` inclui `router` e chama `telas(db, out)` antes de
`montar_grupos` (abas em `_fin_grupos.GRUPOS`). DDL idempotente em `_ensure`.
"""

from __future__ import annotations

import io
import re
from datetime import date
from decimal import Decimal

from fastapi import APIRouter, Body, Depends, HTTPException, Query
from fastapi.responses import Response
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

from core.auth.dependencies import CurrentActiveUser, require_permission
from core.database import get_db

#: ANTES do import do data_controller (o ciclo de import fecha com o router pronto) — igual F11.
router = APIRouter()

from modules.operacional.controllers.redesign_data_controller import (  # noqa: E402
    _helpers,
    b,
    brl,
    doc,
    t,
)

_ND = "#0F1B3A"
_ACT = "/api/v1/redesign/action/"
_GATE = [Depends(require_permission("module:financeiro"))]
#: "empresa" do financeiro da casa — o mesmo id que os 369 pagáveis usam (não é condomínio cliente).
_EMPRESA = "a1b2c3d4-e5f6-7890-abcd-ef1234567890"

IDS = (
    "formas-pagamento",
    "forma-pagamento-nova",
    "centros-custo",
    "centro-custo-novo",
    "centro-custo-categoria",
    "relatorio-financeiro",
)

#: tipos do cadastro — os MESMOS valores de `PaymentMethodType` (financial/models/payment_method.py),
#: para o `payment_type` gravado aqui servir ao resto do módulo sem tradução.
TIPOS_FORMA = (
    ("pix", "PIX"),
    ("boleto", "Boleto"),
    ("ted", "TED"),
    ("dinheiro", "Dinheiro"),
    ("cartao_credito", "Cartão de crédito"),
    ("cartao_debito", "Cartão de débito"),
    ("debito_automatico", "Débito automático"),
    ("transferencia", "Transferência"),
    ("outro", "Outro"),
)

#: Semente = a lista canônica acima. O oráculo prova que ela COBRE 100% dos valores distintos de
#: forma de pagamento que já existem no banco (`inter_payments.payment_type`, `payable_payments.
#: payment_method_name`, `receivable_payments.payment_method_name`, `commission_payments.
#: payment_method`) — medido em 24/09: {pix: 31, boleto: 16} e nada mais.
_SEED_FORMAS = [(c, n, c == "boleto") for c, n in TIPOS_FORMA]

CENTER_TYPES = (
    ("OPERATIONAL", "Operacional"),
    ("ADMINISTRATIVE", "Administrativo"),
    ("COMMERCIAL", "Comercial"),
    ("SUPPORT", "Apoio"),
    ("PROJECT", "Projeto"),
    ("DEPARTMENT", "Departamento"),
    ("BRANCH", "Filial"),
    ("OTHER", "Outro"),
)

_DDL = [
    # 1) cadastro fiscal completo do código de serviço (LC 214/2025 — IBS/CBS). NULL = sem fonte.
    "ALTER TABLE fin_codigos_servico ADD COLUMN IF NOT EXISTS nbs text",
    "ALTER TABLE fin_codigos_servico ADD COLUMN IF NOT EXISTS cst_iss text",
    "ALTER TABLE fin_codigos_servico ADD COLUMN IF NOT EXISTS cst_pis text",
    "ALTER TABLE fin_codigos_servico ADD COLUMN IF NOT EXISTS cst_cofins text",
    "ALTER TABLE fin_codigos_servico ADD COLUMN IF NOT EXISTS classificacao_tributaria text",
    "ALTER TABLE fin_codigos_servico ADD COLUMN IF NOT EXISTS incide_ibs boolean",
    "ALTER TABLE fin_codigos_servico ADD COLUMN IF NOT EXISTS incide_cbs boolean",
    "ALTER TABLE fin_codigos_servico ADD COLUMN IF NOT EXISTS aliquota_ibs numeric(5,2)",
    "ALTER TABLE fin_codigos_servico ADD COLUMN IF NOT EXISTS aliquota_cbs numeric(5,2)",
    "ALTER TABLE fin_codigos_servico ADD COLUMN IF NOT EXISTS origem_regra text",
    # 2) limite de valor por condição de pagamento (regra que o servidor do DGX impõe)
    "ALTER TABLE fin_condicoes_pagamento ADD COLUMN IF NOT EXISTS limite_valor numeric(14,2)",
    # 3) categoria do extrato → centro de custo formal (sem migrar dado: o extrato não muda)
    """CREATE TABLE IF NOT EXISTS fin_cost_center_categorias (
        cost_center_id uuid NOT NULL, categoria text NOT NULL PRIMARY KEY,
        created_at timestamptz DEFAULT now())""",
    "CREATE INDEX IF NOT EXISTS ix_cc_categorias_centro ON fin_cost_center_categorias (cost_center_id)",
    # 4) forma de pagamento no recebível (o pagável já tem payment_method_id; o recebível não tinha)
    "ALTER TABLE receivable_accounts ADD COLUMN IF NOT EXISTS payment_method_id uuid",
]


async def _ensure(db: AsyncSession) -> None:
    for sql in _DDL:
        await db.execute(text(sql))
    for cod, nome, boleto in _SEED_FORMAS:
        await db.execute(
            text(
                "INSERT INTO payment_methods (id, condominio_id, code, name, payment_type, status, "
                "  requires_bank_account, display_order, is_default, ativo, created_at, updated_at) "
                "SELECT gen_random_uuid(), CAST(:e AS uuid), CAST(:c AS varchar), CAST(:n AS varchar), "
                "  CAST(:c AS varchar), 'ativo', :rb, 0, false, true, now(), now() "
                " WHERE NOT EXISTS (SELECT 1 FROM payment_methods WHERE code = CAST(:c AS varchar))"
            ),
            {"e": _EMPRESA, "c": cod, "n": nome, "rb": bool(boleto)},
        )
    await db.commit()


# ── regras puras (o oráculo importa estas) ──────────────────────────────────────────────────
def ancestrais(pais: dict[str, str | None], cid: str, limite: int = 50) -> tuple[list[str], bool]:
    """(cadeia de pais de `cid`, do pai à raiz; houve ciclo). Para ao repetir um nó — nunca laça."""
    out: list[str] = []
    visto = {cid}
    p = pais.get(cid)
    while p and p not in visto and len(out) < limite:
        out.append(p)
        visto.add(p)
        p = pais.get(p)
    return out, p is not None and p in visto


def arvore_valida(pais: dict[str, str | None]) -> list[str]:
    """Problemas da árvore: pai inexistente (órfão) ou ciclo. Lista vazia = sadia."""
    problemas = []
    for cid, pai in pais.items():
        if pai is None:
            continue
        if pai not in pais:
            problemas.append(f"órfão: centro {cid} aponta para o pai {pai}, que não existe")
            continue
        _cadeia, laco = ancestrais(pais, cid)
        if laco:
            problemas.append(f"ciclo: centro {cid} é ancestral de si mesmo")
    return problemas


def _d(v, rotulo: str) -> date:
    s = str(v or "").strip()
    if re.fullmatch(r"\d{2}/\d{2}/\d{4}", s):
        dd, mm, aa = s.split("/")
        s = f"{aa}-{mm}-{dd}"
    try:
        return date.fromisoformat(s)
    except ValueError:
        raise HTTPException(status_code=400, detail=f"{rotulo}: use DD/MM/AAAA.")


def _br(d: date | None) -> str:
    return d.strftime("%d/%m/%Y") if d else "—"


# ── relatórios ──────────────────────────────────────────────────────────────────────────────
#: tipo → (título, FROM/WHERE base, {agrupamento: (rótulo, expressão-chave)}, agrupamento padrão)
_REL = {
    "contas-pagar": (
        "Contas a pagar",
        "FROM payable_accounts WHERE due_date BETWEEN :de AND :ate "
        "  AND coalesce(status,'') <> 'cancelada' AND coalesce(ativo, true)",
        {
            "fornecedor": (
                "Fornecedor",
                "coalesce(nullif(supplier_name,''), nullif(fornecedor_nome,''), '(sem fornecedor)')",
            ),
            "vencimento": ("Vencimento", "to_char(due_date, 'DD/MM/YYYY')"),
            "centro": ("Centro de custo", "coalesce(nullif(cost_center,''), '(sem centro)')"),
            "categoria": ("Categoria", "coalesce(nullif(categoria,''), '(sem categoria)')"),
        },
        "fornecedor",
    ),
    "contas-receber": (
        "Contas a receber",
        "FROM receivable_accounts WHERE due_date BETWEEN :de AND :ate "
        "  AND coalesce(status,'') <> 'cancelada' AND coalesce(ativo, true)",
        {
            "cliente": ("Cliente", "coalesce(nullif(customer_name,''), '(sem cliente)')"),
            "vencimento": ("Vencimento", "to_char(due_date, 'DD/MM/YYYY')"),
            "status": ("Situação", "coalesce(nullif(status,''), '(sem situação)')"),
        },
        "cliente",
    ),
    "fluxo-caixa": (
        "Fluxo de caixa (extrato)",
        "FROM bank_transactions WHERE transaction_date BETWEEN :de AND :ate",
        {
            "dia": ("Dia", "to_char(transaction_date, 'DD/MM/YYYY')"),
            "semana": ("Semana de", "to_char(date_trunc('week', transaction_date), 'DD/MM/YYYY')"),
            "mes": ("Mês", "to_char(transaction_date, 'MM/YYYY')"),
        },
        "mes",
    ),
    "contas-fixas": (
        "Contas fixas",
        "FROM financial_custos_recorrentes WHERE coalesce(ativo, true) AND encerrado_em IS NULL",
        {"categoria": ("Categoria", "coalesce(nullif(categoria,''), '(sem categoria)')")},
        "categoria",
    ),
}
#: coluna de rótulo e de valor por tipo (cada tabela chama o seu de um jeito)
_REL_CAMPOS = {
    "contas-pagar": ("coalesce(description, '(sem descrição)')", "coalesce(gross_value, 0)"),
    "contas-receber": ("coalesce(description, '(sem descrição)')", "coalesce(gross_value, 0)"),
    "fluxo-caixa": ("coalesce(description, '(sem descrição)')", "coalesce(amount, 0)"),
    "contas-fixas": ("coalesce(descricao, '(sem descrição)')", "coalesce(valor, 0)"),
}


async def dados_relatorio(db: AsyncSession, tipo: str, de: date, ate: date, agrupar: str = "") -> dict:
    """Linhas do relatório agrupadas. O oráculo reconta o total por SQL próprio."""
    if tipo not in _REL:
        raise HTTPException(status_code=400, detail=f"Relatório desconhecido: {tipo}.")
    titulo, base, agrs, padrao = _REL[tipo]
    ag = (agrupar or padrao).strip() or padrao
    if ag not in agrs:
        raise HTTPException(status_code=400, detail=f"Agrupamento inválido para {tipo}: {ag}.")
    rot_ag, expr = agrs[ag]
    col_rot, col_val = _REL_CAMPOS[tipo]
    sql = f"SELECT {expr} AS chave, {col_rot} AS rotulo, {col_val} AS valor {base} ORDER BY 1, 2"
    params = {"de": de, "ate": ate} if ":de" in base else {}
    linhas = (await db.execute(text(sql), params)).fetchall()
    grupos: list[dict] = []
    atual = None
    for chave, rotulo, valor in linhas:
        if atual is None or atual["nome"] != str(chave):
            atual = {"nome": str(chave), "linhas": [], "total": Decimal("0")}
            grupos.append(atual)
        v = Decimal(str(valor or 0))
        atual["linhas"].append((str(rotulo)[:70], v))
        atual["total"] += v
    return {
        "tipo": tipo,
        "titulo": titulo,
        "agrupar": ag,
        "rotulo_agrupamento": rot_ag,
        "de": de,
        "ate": ate,
        "grupos": grupos,
        "total": sum((g["total"] for g in grupos), Decimal("0")),
        "qtd": sum(len(g["linhas"]) for g in grupos),
    }


def _sub(d: dict) -> str:
    return f"{_br(d['de'])} a {_br(d['ate'])} · por {d['rotulo_agrupamento'].lower()} · {d['qtd']} lançamento(s)"


def relatorio_pdf(d: dict) -> bytes:
    """PDF timbrado — reusa o gerador da casa (`relatorio_financeiro_pdf`), que já É a marca."""
    from modules.financial.services.relatorio_financeiro_pdf import gerar_relatorio_pdf

    secoes = [
        {
            "titulo": f"{d['rotulo_agrupamento']}: {g['nome']}",
            "linhas": [
                *[(r, float(v)) for r, v in g["linhas"]],
                (f"Subtotal — {g['nome']}", float(g["total"]), True),
            ],
        }
        for g in d["grupos"]
    ] or [{"titulo": "Sem lançamentos no período", "linhas": [("(nada a mostrar)", "")]}]
    secoes.append({"titulo": "Total", "linhas": [("TOTAL DO PERÍODO", float(d["total"]), True)]})
    return gerar_relatorio_pdf(d["titulo"], _sub(d), secoes)


def relatorio_xlsx(d: dict) -> bytes:
    """Excel por openpyxl (3.1.5 no container) — uma linha por lançamento + subtotal e total."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = d["titulo"][:31]
    neg = Font(bold=True)
    ws.append([d["titulo"]])
    ws["A1"].font = neg
    ws.append([_sub(d)])
    ws.append([])
    ws.append([d["rotulo_agrupamento"], "Descrição", "Valor (R$)"])
    for c in ws[4]:
        c.font = neg
    for g in d["grupos"]:
        for rot, val in g["linhas"]:
            ws.append([g["nome"], rot, float(val)])
        ws.append([g["nome"], f"Subtotal — {g['nome']}", float(g["total"])])
        for c in ws[ws.max_row]:
            c.font = neg
    ws.append([])
    ws.append(["", "TOTAL DO PERÍODO", float(d["total"])])
    for c in ws[ws.max_row]:
        c.font = neg
    for col, larg in (("A", 28), ("B", 62), ("C", 16)):
        ws.column_dimensions[col].width = larg
    for linha in ws.iter_rows(min_row=4, min_col=3, max_col=3):
        for c in linha:
            c.number_format = "#,##0.00"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@router.get("/relatorio-financeiro/{tipo}", dependencies=_GATE)
async def rd_relatorio_financeiro(
    tipo: str,
    current_user: CurrentActiveUser,
    de: str = Query("", description="DD/MM/AAAA ou AAAA-MM-DD"),
    ate: str = Query(""),
    agrupar: str = Query(""),
    fmt: str = Query("pdf"),
    db: AsyncSession = Depends(get_db),
) -> Response:
    hoje = date.today()
    d1 = _d(de, "De") if de else hoje.replace(day=1)
    d2 = _d(ate, "Até") if ate else hoje
    if d2 < d1:
        raise HTTPException(status_code=400, detail="Período invertido: «até» é antes de «de».")
    dados = await dados_relatorio(db, tipo, d1, d2, agrupar)
    nome = f"{tipo}_{d1:%Y%m%d}_{d2:%Y%m%d}"
    if fmt == "xlsx":
        return Response(
            content=relatorio_xlsx(dados),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{nome}.xlsx"'},
        )
    return Response(
        content=relatorio_pdf(dados),
        media_type="application/pdf",
        headers={"Content-Disposition": f'inline; filename="{nome}.pdf"'},
    )


# ── telas ───────────────────────────────────────────────────────────────────────────────────
async def telas(db, out: dict | None = None) -> dict:
    await _ensure(db)
    mine, safe, tbl = _helpers(db)
    out = out if out is not None else {}
    hoje = date.today()
    prim = hoje.replace(day=1)

    # 1) Formas de pagamento ──────────────────────────────────────────────────────────────
    contas = (
        await db.execute(
            text(
                "SELECT id::text, coalesce(bank_name,'') || ' · ' || coalesce(account_number,'') "
                "  FROM bank_accounts ORDER BY bank_name"
            )
        )
    ).fetchall()
    _opt_conta = [{"value": "", "label": "— nenhuma —"}] + [{"value": c[0], "label": c[1][:44]} for c in contas]
    _campos_forma = lambda r=None: [  # noqa: E731
        {
            "key": "code",
            "label": "Código*",
            "type": "text",
            "span": "span 1",
            "ph": "pix",
            **({"value": r[1] or ""} if r else {}),
        },
        {
            "key": "name",
            "label": "Nome*",
            "type": "text",
            "span": "span 1",
            "ph": "PIX",
            **({"value": r[2] or ""} if r else {}),
        },
        {
            "key": "payment_type",
            "label": "Tipo*",
            "type": "select",
            "span": "span 1",
            "options": [{"value": v, "label": lb} for v, lb in TIPOS_FORMA],
            **({"value": r[3] or ""} if r else {}),
        },
        {
            "key": "bank_account_id",
            "label": "Conta bancária (TED, débito automático, boleto)",
            "type": "select",
            "span": "span 1",
            "options": _opt_conta,
            **({"value": r[4] or ""} if r else {}),
        },
    ]
    await safe(
        "formas-pagamento",
        tbl(
            "Formas de pagamento",
            "Cadastro que os títulos apontam — `payment_method_id` já existia em pagáveis e baixas e estava sem "
            "nada para apontar (tabela com 0 linhas). «Gera boleto» = tipo boleto. Inativar não apaga histórico.",
            "Nova forma",
            ["Código", "Nome", "Tipo", "Gera boleto", "Conta bancária", "Em uso", "Ativa"],
            "0.8fr 1.4fr 1.2fr 0.8fr 1.6fr 0.6fr 0.6fr",
            "SELECT m.id::text, m.code, m.name, m.payment_type, m.bank_account_id::text, coalesce(m.ativo,true), "
            "       coalesce(ba.bank_name,'—'), "
            "       (SELECT count(*) FROM payable_accounts p WHERE p.payment_method_id = m.id) "
            "     + (SELECT count(*) FROM payable_payments p WHERE p.payment_method_id = m.id) "
            "     + (SELECT count(*) FROM receivable_payments p WHERE p.payment_method_id = m.id) "
            "  FROM payment_methods m LEFT JOIN bank_accounts ba ON ba.id = m.bank_account_id "
            " ORDER BY coalesce(m.ativo,true) DESC, m.code",
            lambda r: [
                t(r[1] or "—", 600, _ND),
                t(r[2] or "—", 600),
                b(dict(TIPOS_FORMA).get(r[3] or "", r[3] or "—"), "info"),
                b("sim", "ok") if r[3] == "boleto" else t("não"),
                t(str(r[6])[:30]),
                b(str(r[7]), "ok" if r[7] else "mut"),
                b("Ativa", "ok") if r[5] else b("Inativa", "mut"),
            ],
            editfn=lambda r: {
                "endpoint": f"{_ACT}forma-pagamento-salvar",
                "method": "POST",
                "btnLabel": "Editar",
                "submitLabel": "Salvar",
                "okMsg": "Forma salva. Recarregue.",
                "fields": [
                    {"key": "id", "label": "id", "type": "text", "value": r[0], "span": "span 1"},
                    *_campos_forma(r),
                ],
            },
            actionsfn=lambda r: [
                {
                    "title": f"{'Inativar' if r[5] else 'Reativar'} «{r[2]}»",
                    "endpoint": f"{_ACT}forma-pagamento-salvar",
                    "method": "POST",
                    "btnLabel": "Inativar" if r[5] else "Reativar",
                    "btnStyle": "outline",
                    "submitLabel": "Confirmar",
                    "okMsg": "Feito. Recarregue.",
                    "fields": [
                        {"key": "id", "label": "id", "type": "text", "value": r[0], "span": "span 1"},
                        {
                            "key": "ativo",
                            "label": "ativo",
                            "type": "text",
                            "value": "0" if r[5] else "1",
                            "span": "span 1",
                        },
                    ],
                }
            ],
        ),
    )
    if isinstance(mine.get("formas-pagamento"), dict):
        mine["formas-pagamento"]["ctaTo"] = "forma-pagamento-nova"
    mine["forma-pagamento-nova"] = {
        "title": "Nova forma de pagamento",
        "sub": "Código curto (o que o fornecedor e as baixas apontam) + tipo. Conta bancária só quando a forma sai de uma conta.",
        "cta": "Salvar",
        "type": "form",
        "submit": {"endpoint": f"{_ACT}forma-pagamento-salvar", "okMsg": "Forma de pagamento criada."},
        "fields": _campos_forma(),
    }

    # o select entra nos forms de «Registrar conta» (o F11 já trocou o endpoint deles para
    # payable-condicao/receivable-condicao — é lá que o id é gravado no título criado).
    formas = (
        await db.execute(text("SELECT id::text, name FROM payment_methods WHERE coalesce(ativo,true) ORDER BY code"))
    ).fetchall()
    _opt_forma = [{"value": "", "label": "— não informar —"}] + [{"value": f[0], "label": f[1]} for f in formas]
    for fid in ("registrar-conta-pagar", "registrar-conta-receber"):
        scr = out.get(fid)
        if (
            isinstance(scr, dict)
            and scr.get("type") == "form"
            and not any(f.get("key") == "forma_pagamento_id" for f in scr.get("fields", []))
        ):
            scr["fields"].append(
                {
                    "key": "forma_pagamento_id",
                    "label": "Forma de pagamento",
                    "type": "select",
                    "span": "span 1",
                    "options": _opt_forma,
                }
            )

    # 2) Centros de custo em árvore ───────────────────────────────────────────────────────
    ccs = (
        await db.execute(
            text(
                "SELECT c.id::text, c.parent_id::text, c.code, c.name, c.center_type::text, c.level, "
                "       coalesce(c.full_path, c.name), coalesce(c.active, true), "
                "       (SELECT count(*) FROM fin_cost_center_categorias k WHERE k.cost_center_id = c.id), "
                "       (SELECT count(*) FROM fin_cost_centers f WHERE f.parent_id = c.id) "
                "  FROM fin_cost_centers c ORDER BY coalesce(c.full_path, c.name)"
            )
        )
    ).fetchall()
    cats_por_centro = {
        r[0]: r[1]
        for r in (
            await db.execute(
                text(
                    "SELECT cost_center_id::text, string_agg(categoria, ', ' ORDER BY categoria) "
                    "  FROM fin_cost_center_categorias GROUP BY 1"
                )
            )
        ).fetchall()
    }
    mine["centros-custo"] = {
        "title": "Centros de custo (árvore)",
        "sub": f"{len(ccs)} centro(s). A hierarquia usa `parent_id`, que a tabela já tinha. Cada centro recebe uma ou "
        "mais CATEGORIAS do extrato — é assim que a análise orçamentária passa a falar «por centro» sem migrar "
        "um lançamento sequer.",
        "cta": "Novo centro",
        "ctaTo": "centro-custo-novo",
        "type": "table",
        "searchHint": "Buscar centro…",
        "grid": "0.7fr 2.2fr 1.1fr 1.4fr 0.6fr 0.6fr",
        "cols": ["Código", "Centro", "Tipo", "Categorias do extrato", "Filhos", "Ativo"],
        "rows": [
            {
                "cells": [
                    t(c[2] or "—", 600, _ND),
                    t(("— " * max(0, int(c[5] or 1) - 1)) + (c[3] or "—"), 600 if (c[5] or 1) == 1 else 400),
                    b(dict(CENTER_TYPES).get(c[4] or "", c[4] or "—"), "info"),
                    t((cats_por_centro.get(c[0]) or "—")[:44]),
                    b(str(c[9]), "ok" if c[9] else "mut"),
                    b("Ativo", "ok") if c[7] else b("Inativo", "mut"),
                ]
            }
            for c in ccs
        ]
        or [
            {
                "cells": [
                    t("Nenhum centro cadastrado", 500, "#94A3B8"),
                    t("—"),
                    t("—"),
                    t("—"),
                    b("—", "mut"),
                    b("—", "mut"),
                ]
            }
        ],
    }
    _opt_pai = [{"value": "", "label": "— raiz (sem pai) —"}] + [
        {"value": c[0], "label": f"{c[2]} · {c[3]}"[:48]} for c in ccs if c[7]
    ]
    mine["centro-custo-novo"] = {
        "title": "Novo centro de custo",
        "sub": "Código único + nome. Escolher um pai cria o nível abaixo dele; pai que fecharia ciclo é recusado.",
        "cta": "Salvar",
        "type": "form",
        "submit": {"endpoint": f"{_ACT}centro-custo-salvar", "okMsg": "Centro de custo criado."},
        "fields": [
            {"key": "code", "label": "Código*", "type": "text", "span": "span 1", "ph": "01.02"},
            {"key": "name", "label": "Nome*", "type": "text", "span": "span 2", "ph": "Ex.: Portaria — Michelangelo"},
            {
                "key": "center_type",
                "label": "Tipo*",
                "type": "select",
                "span": "span 1",
                "value": "OPERATIONAL",
                "options": [{"value": v, "label": lb} for v, lb in CENTER_TYPES],
            },
            {"key": "parent_id", "label": "Centro pai", "type": "select", "span": "span 1", "options": _opt_pai},
        ],
    }
    cats = [
        r[0]
        for r in (
            await db.execute(
                text(
                    "SELECT DISTINCT category FROM bank_transactions "
                    " WHERE category IS NOT NULL AND category <> '' ORDER BY 1"
                )
            )
        ).fetchall()
    ]
    ligadas = {r[0] for r in (await db.execute(text("SELECT categoria FROM fin_cost_center_categorias"))).fetchall()}
    mine["centro-custo-categoria"] = {
        "title": "Ligar categoria do extrato a um centro",
        "sub": f"{len(cats)} categoria(s) no extrato · {len(ligadas)} já ligada(s). Uma categoria pertence a UM centro. "
        "Sem isto, a análise orçamentária só fala por categoria — que é o que a casa usa hoje.",
        "cta": "Ligar",
        "type": "form",
        "submit": {"endpoint": f"{_ACT}centro-custo-categoria", "okMsg": "Categoria ligada ao centro."},
        "fields": [
            {
                "key": "categoria",
                "label": "Categoria do extrato*",
                "type": "select",
                "span": "span 1",
                "options": [{"value": c, "label": c + (" (já ligada)" if c in ligadas else "")} for c in cats]
                or [{"value": "", "label": "— sem categorias no extrato —"}],
            },
            {
                "key": "cost_center_id",
                "label": "Centro de custo* (vazio = desligar)",
                "type": "select",
                "span": "span 1",
                "options": [{"value": "", "label": "— desligar —"}]
                + [{"value": c[0], "label": f"{c[2]} · {c[3]}"[:48]} for c in ccs if c[7]],
            },
        ],
    }

    # 3) Relatórios financeiros em PDF/Excel ──────────────────────────────────────────────
    _url = lambda tp, ag, fm: (  # noqa: E731
        f"/api/v1/redesign/relatorio-financeiro/{tp}?de={prim:%Y-%m-%d}&ate={hoje:%Y-%m-%d}&agrupar={ag}&fmt={fm}"
    )
    mine["relatorio-financeiro"] = {
        "title": "Relatórios financeiros (PDF timbrado / Excel)",
        "sub": "Escolha o relatório, o período e o agrupamento — sai o PDF da marca ou a planilha. "
        f"Os botões do topo são o atalho do mês corrente ({prim:%d/%m} a {hoje:%d/%m}).",
        "cta": "Gerar",
        "type": "form",
        "submit": {
            "endpoint": f"{_ACT}relatorio-financeiro",
            "okMsg": "Relatório gerado — abrindo.",
            "showResult": True,
        },
        "fields": [
            {
                "key": "tipo",
                "label": "Relatório*",
                "type": "select",
                "span": "span 1",
                "value": "contas-pagar",
                "options": [
                    {"value": "contas-pagar", "label": "Contas a pagar"},
                    {"value": "contas-receber", "label": "Contas a receber"},
                    {"value": "fluxo-caixa", "label": "Fluxo de caixa (extrato)"},
                    {"value": "contas-fixas", "label": "Contas fixas"},
                ],
            },
            {
                "key": "agrupar",
                "label": "Agrupar por*",
                "type": "select",
                "span": "span 1",
                "value": "fornecedor",
                "options": [
                    {"value": "fornecedor", "label": "A pagar: fornecedor"},
                    {"value": "vencimento", "label": "A pagar / a receber: vencimento"},
                    {"value": "centro", "label": "A pagar: centro de custo"},
                    {"value": "categoria", "label": "A pagar / contas fixas: categoria"},
                    {"value": "cliente", "label": "A receber: cliente"},
                    {"value": "status", "label": "A receber: situação"},
                    {"value": "dia", "label": "Fluxo de caixa: dia"},
                    {"value": "semana", "label": "Fluxo de caixa: semana"},
                    {"value": "mes", "label": "Fluxo de caixa: mês"},
                ],
            },
            {"key": "de", "label": "De (DD/MM/AAAA)*", "type": "text", "span": "span 1", "value": f"{prim:%d/%m/%Y}"},
            {"key": "ate", "label": "Até (DD/MM/AAAA)*", "type": "text", "span": "span 1", "value": f"{hoje:%d/%m/%Y}"},
            {
                "key": "fmt",
                "label": "Formato*",
                "type": "select",
                "span": "span 1",
                "value": "pdf",
                "options": [{"value": "pdf", "label": "PDF timbrado"}, {"value": "xlsx", "label": "Excel (.xlsx)"}],
            },
        ],
        "docs": [
            doc(
                "A pagar por fornecedor (PDF)", _url("contas-pagar", "fornecedor", "pdf"), fmt="pdf", gate="financeiro"
            ),
            doc(
                "A pagar por fornecedor (Excel)",
                _url("contas-pagar", "fornecedor", "xlsx"),
                fmt="xlsx",
                gate="financeiro",
            ),
            doc("A receber por cliente (PDF)", _url("contas-receber", "cliente", "pdf"), fmt="pdf", gate="financeiro"),
            doc("Fluxo de caixa por dia (PDF)", _url("fluxo-caixa", "dia", "pdf"), fmt="pdf", gate="financeiro"),
            doc("Contas fixas (Excel)", _url("contas-fixas", "categoria", "xlsx"), fmt="xlsx", gate="financeiro"),
        ],
    }

    out.update(mine)
    return out


# ── ações ───────────────────────────────────────────────────────────────────────────────────
@router.post("/action/forma-pagamento-salvar", dependencies=_GATE)
async def rd_forma_pagamento_salvar(
    current_user: CurrentActiveUser, payload: dict = Body(...), db: AsyncSession = Depends(get_db)
) -> dict:
    await _ensure(db)
    fid = str(payload.get("id") or "").strip()
    if fid and "ativo" in payload and "name" not in payload:
        liga = str(payload["ativo"]) == "1"
        await db.execute(
            text("UPDATE payment_methods SET ativo = :a, status = :s, updated_at = now() WHERE id = CAST(:i AS uuid)"),
            {"a": liga, "s": "ativo" if liga else "inativo", "i": fid},
        )
        await db.commit()
        return {"ok": True, "message": "Forma de pagamento atualizada."}
    code = str(payload.get("code") or "").strip().lower()
    nome = str(payload.get("name") or "").strip()
    tipo = str(payload.get("payment_type") or "").strip()
    if not re.fullmatch(r"[a-z0-9_-]{2,20}", code):
        raise HTTPException(status_code=400, detail="Código: 2 a 20 letras minúsculas, dígitos, _ ou -.")
    if len(nome) < 2:
        raise HTTPException(status_code=400, detail="Nome obrigatório.")
    if tipo not in dict(TIPOS_FORMA):
        raise HTTPException(status_code=400, detail=f"Tipo de forma de pagamento inválido: «{tipo}».")
    conta = str(payload.get("bank_account_id") or "").strip() or None
    p = {"c": code, "n": nome, "t": tipo, "b": conta, "e": _EMPRESA}
    if fid:
        await db.execute(
            text(
                "UPDATE payment_methods SET code=CAST(:c AS varchar), name=CAST(:n AS varchar), "
                " payment_type=CAST(:t AS varchar), bank_account_id=CAST(:b AS uuid), "
                " requires_bank_account = (CAST(:t AS varchar) = 'boleto'), updated_at=now() "
                " WHERE id=CAST(:i AS uuid)"
            ),
            {**p, "i": fid},
        )
    else:
        if (
            await db.execute(text("SELECT 1 FROM payment_methods WHERE code = CAST(:c AS varchar)"), {"c": code})
        ).first():
            raise HTTPException(status_code=400, detail=f"Já existe forma com o código «{code}».")
        await db.execute(
            text(
                "INSERT INTO payment_methods (id, condominio_id, code, name, payment_type, status, "
                " requires_bank_account, bank_account_id, display_order, is_default, ativo, created_at, updated_at) "
                "VALUES (gen_random_uuid(), CAST(:e AS uuid), CAST(:c AS varchar), CAST(:n AS varchar), "
                " CAST(:t AS varchar), 'ativo', (CAST(:t AS varchar) = 'boleto'), CAST(:b AS uuid), 0, false, true, "
                " now(), now())"
            ),
            p,
        )
    await db.commit()
    return {"ok": True, "message": f"Forma «{nome}» salva."}


@router.post("/action/centro-custo-salvar", dependencies=_GATE)
async def rd_centro_custo_salvar(
    current_user: CurrentActiveUser, payload: dict = Body(...), db: AsyncSession = Depends(get_db)
) -> dict:
    await _ensure(db)
    code = str(payload.get("code") or "").strip()
    nome = str(payload.get("name") or "").strip()
    tipo = str(payload.get("center_type") or "OPERATIONAL").strip().upper()
    pai = str(payload.get("parent_id") or "").strip() or None
    if not 1 <= len(code) <= 20:
        raise HTTPException(status_code=400, detail="Código: 1 a 20 caracteres.")
    if len(nome) < 2:
        raise HTTPException(status_code=400, detail="Nome obrigatório.")
    if tipo not in dict(CENTER_TYPES):
        raise HTTPException(status_code=400, detail=f"Tipo de centro inválido: «{tipo}».")
    if (await db.execute(text("SELECT 1 FROM fin_cost_centers WHERE code = CAST(:c AS varchar)"), {"c": code})).first():
        raise HTTPException(status_code=400, detail=f"Já existe centro com o código «{code}».")
    nivel, caminho = 1, nome
    if pai:
        p = (
            await db.execute(
                text("SELECT level, coalesce(full_path, name) FROM fin_cost_centers WHERE id = CAST(:i AS uuid)"),
                {"i": pai},
            )
        ).first()
        if not p:
            raise HTTPException(status_code=400, detail="Centro pai inexistente.")
        nivel, caminho = int(p[0] or 1) + 1, f"{p[1]} / {nome}"
        if nivel > 8:
            raise HTTPException(status_code=400, detail="Árvore profunda demais (máximo 8 níveis).")
    await db.execute(
        text(
            "INSERT INTO fin_cost_centers (id, condominio_id, parent_id, code, name, center_type, status, "
            " allocation_method, level, full_path, budget_amount, actual_amount, accepts_entries, is_productive, "
            " is_allocatable, active, created_at, updated_at) "
            "VALUES (gen_random_uuid(), CAST(:e AS uuid), CAST(:p AS uuid), CAST(:c AS varchar), CAST(:n AS varchar), "
            " CAST(:t AS costcentertype), 'ACTIVE', 'DIRECT', :l, CAST(:f AS varchar), 0, 0, true, true, true, true, "
            " now(), now())"
        ),
        {"e": _EMPRESA, "p": pai, "c": code, "n": nome, "t": tipo, "l": nivel, "f": caminho[:255]},
    )
    await db.commit()
    return {"ok": True, "message": f"Centro «{code} · {nome}» criado no nível {nivel}."}


@router.post("/action/centro-custo-categoria", dependencies=_GATE)
async def rd_centro_custo_categoria(
    current_user: CurrentActiveUser, payload: dict = Body(...), db: AsyncSession = Depends(get_db)
) -> dict:
    await _ensure(db)
    cat = str(payload.get("categoria") or "").strip()
    cid = str(payload.get("cost_center_id") or "").strip()
    if not cat:
        raise HTTPException(status_code=400, detail="Categoria obrigatória.")
    if not cid:
        await db.execute(text("DELETE FROM fin_cost_center_categorias WHERE categoria = :c"), {"c": cat})
        await db.commit()
        return {"ok": True, "message": f"Categoria «{cat}» desligada de qualquer centro."}
    centro = (
        await db.execute(text("SELECT code, name FROM fin_cost_centers WHERE id = CAST(:i AS uuid)"), {"i": cid})
    ).first()
    if not centro:
        raise HTTPException(status_code=400, detail="Centro de custo inexistente.")
    await db.execute(
        text(
            "INSERT INTO fin_cost_center_categorias (cost_center_id, categoria) VALUES (CAST(:i AS uuid), :c) "
            "ON CONFLICT (categoria) DO UPDATE SET cost_center_id = EXCLUDED.cost_center_id"
        ),
        {"i": cid, "c": cat},
    )
    await db.commit()
    return {"ok": True, "message": f"«{cat}» agora pertence a {centro[0]} · {centro[1]}."}


@router.post("/action/relatorio-financeiro", dependencies=_GATE)
async def rd_relatorio_financeiro_acao(
    current_user: CurrentActiveUser, payload: dict = Body(...), db: AsyncSession = Depends(get_db)
) -> dict:
    """Monta a URL do relatório com os filtros escolhidos e devolve `doc` — o form abre direto."""
    tipo = str(payload.get("tipo") or "contas-pagar").strip()
    if tipo not in _REL:
        raise HTTPException(status_code=400, detail=f"Relatório desconhecido: {tipo}.")
    hoje = date.today()
    d1 = _d(payload.get("de"), "De") if str(payload.get("de") or "").strip() else hoje.replace(day=1)
    d2 = _d(payload.get("ate"), "Até") if str(payload.get("ate") or "").strip() else hoje
    if d2 < d1:
        raise HTTPException(status_code=400, detail="Período invertido: «até» é antes de «de».")
    _titulo, _base, agrs, padrao = _REL[tipo]
    ag = str(payload.get("agrupar") or "").strip()
    if ag not in agrs:
        ag = padrao
    fmt = "xlsx" if str(payload.get("fmt") or "pdf").strip() == "xlsx" else "pdf"
    dados = await dados_relatorio(db, tipo, d1, d2, ag)
    url = f"/api/v1/redesign/relatorio-financeiro/{tipo}?de={d1:%Y-%m-%d}&ate={d2:%Y-%m-%d}&agrupar={ag}&fmt={fmt}"
    return {
        "ok": True,
        "doc": {"label": dados["titulo"], "url": url, "fmt": fmt, "mode": "blob"},
        "relatorio": dados["titulo"],
        "periodo": f"{_br(d1)} a {_br(d2)}",
        "agrupado_por": dados["rotulo_agrupamento"],
        "grupos": len(dados["grupos"]),
        "lancamentos": dados["qtd"],
        "total": brl(dados["total"]),
        "message": (
            f"{dados['titulo']} · {dados['qtd']} lançamento(s) em {len(dados['grupos'])} grupo(s) "
            f"· total {brl(dados['total'])}."
        ),
    }
