"""CFO IA — endpoints (Conecta Mais). Consultor financeiro ancorado nos números reais do ERP.

A IA assiste; o gestor/contador decide. Toda resposta traz disclaimer e sinal de escalonamento.
Espelha o Consultor Jurídico (chat + anexo + histórico), aplicado ao domínio financeiro.
"""

from __future__ import annotations

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, UploadFile
from pydantic import BaseModel, Field
from sqlalchemy.ext.asyncio import AsyncSession

from core.auth.dependencies import get_current_active_user
from core.database import get_db

from . import cfo_service as svc

router = APIRouter(prefix="/financial/cfo", tags=["Financeiro - CFO IA"])


class PerguntaIn(BaseModel):
    area: str = Field(..., description="Lente: fluxo_caixa | resultado | tributos | estrategico",
                      examples=["fluxo_caixa"])
    pergunta: str = Field(..., min_length=3, examples=["Meu caixa cobre a folha deste mês?"])


class ConsultaOut(BaseModel):
    resposta: str
    escalonar: bool = False
    disclaimer: str
    id: int | None = None
    indisponivel: bool = False
    panorama: dict | None = None


@router.get("/panorama", summary="Fotografia financeira real do ERP agora")
async def obter_panorama(
    current_user=Depends(get_current_active_user),
    db: AsyncSession = Depends(get_db),
) -> dict:
    return await svc.panorama(db)


@router.get("/adimplencia-clientes", summary="MRR contratado x recebido no banco, por cliente")
async def adimplencia_clientes(
    current_user=Depends(get_current_active_user),
    db: AsyncSession = Depends(get_db),
):
    return await svc.adimplencia_clientes(db)


@router.get("/recebido-por-cliente", summary="Quanto cada cliente pagou de verdade no banco (recebimentos identificados)")
async def recebido_por_cliente(
    current_user=Depends(get_current_active_user),
    db: AsyncSession = Depends(get_db),
):
    from sqlalchemy import text as _t
    rows = await db.execute(_t("""
        SELECT COALESCE(NULLIF(trim(counterparty_name),''),'(sem nome)') cliente,
               count(*) qtd, round(sum(amount)::numeric,2) total,
               min(transaction_date) primeiro, max(transaction_date) ultimo
        FROM bank_transactions
        WHERE justificativa_categoria='recebimento_cliente' AND amount > 0
        GROUP BY 1 ORDER BY 3 DESC
    """))
    itens = [{"cliente": r[0], "qtd": int(r[1]), "total": float(r[2]),
              "primeiro": str(r[3]), "ultimo": str(r[4])} for r in rows.all()]
    return {"clientes": itens, "total_recebido": round(sum(i["total"] for i in itens), 2),
            "fonte": "recebimentos identificados no extrato do Inter (PIX/boleto com nome do cliente)"}


@router.get("/projecao-caixa", summary="Projeção de caixa recorrente (saldo + MRR − custos mensais)")
async def projecao_caixa(
    meses: int = 6,
    current_user=Depends(get_current_active_user),
    db: AsyncSession = Depends(get_db),
) -> dict:
    return await svc.projecao_caixa(db, meses=max(1, min(meses, 24)))


@router.get("/previsao-custos", summary="Previsibilidade de custos mensais (folha, tributos, diaristas, fornecedores)")
async def previsao_custos(
    current_user=Depends(get_current_active_user),
    db: AsyncSession = Depends(get_db),
) -> dict:
    return await svc.previsao_custos_mensais(db)


@router.get("/custos-recorrentes", summary="Lista custos recorrentes registrados (tributos, parcelamentos, acordos, fixos)")
async def custos_listar(current_user=Depends(get_current_active_user), db: AsyncSession = Depends(get_db)):
    return {"custos": await svc.custos_recorrentes_listar(db)}


class CustoIn(BaseModel):
    categoria: str = Field(..., description="tributo | parcelamento | acordo | fixo | fornecedor")
    descricao: str
    valor: float
    dia_vencimento: int | None = None
    parcelas_total: int | None = None
    parcelas_pagas: int = 0
    observacao: str | None = None


@router.post("/custos-recorrentes", summary="Registra um custo recorrente (tributo, parcelamento, acordo, custo fixo)")
async def custo_criar(body: CustoIn, current_user=Depends(get_current_active_user), db: AsyncSession = Depends(get_db)):
    return await svc.custo_recorrente_criar(
        db, categoria=body.categoria, descricao=body.descricao, valor=body.valor,
        dia_vencimento=body.dia_vencimento, parcelas_total=body.parcelas_total,
        parcelas_pagas=body.parcelas_pagas, observacao=body.observacao,
        user_id=str(getattr(current_user, "id", None)))


@router.delete("/custos-recorrentes/{custo_id}", summary="Remove (inativa) um custo recorrente")
async def custo_remover(custo_id: int, current_user=Depends(get_current_active_user), db: AsyncSession = Depends(get_db)):
    return await svc.custo_recorrente_remover(db, custo_id)


@router.post("/perguntar", summary="Pergunta ao CFO IA (ancorado nos números reais)",
             response_model=ConsultaOut)
async def perguntar(
    body: PerguntaIn,
    current_user=Depends(get_current_active_user),
    db: AsyncSession = Depends(get_db),
) -> ConsultaOut:
    area = (body.area or "").strip().lower()
    if area not in svc.AREAS_VALIDAS:
        raise HTTPException(status_code=422, detail=f"Área inválida '{body.area}'. Válidas: {', '.join(svc.AREAS_VALIDAS)}.")
    try:
        resultado = await svc.consultar(
            db=db, area=area, pergunta=body.pergunta,
            user_id=str(getattr(current_user, "id", None)))
    except ValueError as e:
        raise HTTPException(status_code=422, detail=str(e)) from e
    return ConsultaOut(**resultado)


def _extrair_texto_arquivo(nome: str, data: bytes) -> str:
    n = (nome or "").lower()
    try:
        if n.endswith(".pdf"):
            import fitz
            doc = fitz.open(stream=data, filetype="pdf")
            txt = "\n".join(p.get_text() for p in doc)
            doc.close()
            return txt.strip()
        if n.endswith(".docx"):
            import io
            from docx import Document
            d = Document(io.BytesIO(data))
            return "\n".join(p.text for p in d.paragraphs).strip()
        if n.endswith((".csv", ".xls", ".xlsx")):
            try:
                import io
                import openpyxl
                wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
                linhas = []
                for ws in wb.worksheets:
                    for row in ws.iter_rows(values_only=True):
                        linhas.append("\t".join("" if c is None else str(c) for c in row))
                        if len(linhas) > 400:
                            break
                return "\n".join(linhas).strip()
            except Exception:
                return data.decode("utf-8", "ignore").strip()
        return data.decode("utf-8", "ignore").strip()
    except Exception:  # noqa: BLE001
        return data.decode("utf-8", "ignore").strip()


@router.post("/perguntar-arquivo", summary="Consulta ao CFO analisando um anexo (PDF/DOCX/planilha/TXT)")
async def perguntar_arquivo(
    arquivo: UploadFile = File(...),
    area: str = Form("fluxo_caixa"),
    pergunta: str = Form(""),
    current_user=Depends(get_current_active_user),
    db: AsyncSession = Depends(get_db),
):
    area_n = (area or "").strip().lower()
    if area_n not in svc.AREAS_VALIDAS:
        raise HTTPException(status_code=422, detail=f"Área inválida '{area}'.")
    data = await arquivo.read()
    texto = _extrair_texto_arquivo(arquivo.filename or "", data)
    if len(texto.strip()) < 15:
        raise HTTPException(status_code=422, detail="Não foi possível extrair texto do arquivo.")
    pergunta_final = (pergunta or "").strip() or "Analise este documento financeiro e aponte o que é relevante, riscos e recomendações."
    try:
        resultado = await svc.consultar(
            db=db, area=area_n, pergunta=pergunta_final,
            user_id=str(getattr(current_user, "id", None)),
            anexo_texto=texto, anexo_nome=arquivo.filename)
    except ValueError as e:
        raise HTTPException(status_code=422, detail=str(e)) from e
    return resultado


@router.get("/historico", summary="Histórico de consultas ao CFO IA")
async def historico(
    current_user=Depends(get_current_active_user),
    area: str | None = Query(default=None),
    limit: int = Query(default=50, ge=1, le=200),
    db: AsyncSession = Depends(get_db),
) -> dict:
    if area is not None and area.strip().lower() not in svc.AREAS_VALIDAS:
        raise HTTPException(status_code=422, detail=f"Área inválida '{area}'.")
    consultas = await svc.listar_consultas(db=db, area=area, limit=limit)
    return {"total": len(consultas), "consultas": consultas}
