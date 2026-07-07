"""Gerador de APRESENTAÇÕES padrão Conecta PRO (slides .pptx + .pdf).

Mesma identidade dos documentos padrão-ouro (pdf_branding): laranja #F97316, azul #1E3A5F,
rodapé de marca fixo, capa com faixa, fechamento assinado pelo CEO. O cowork monta a estrutura
(JSON) e este módulo renderiza — para que TODA apresentação saia no padrão, não no visual do editor.

Estrutura de entrada (dict):
{
  "titulo": "Vídeo Monitoramento Inteligente",
  "subtitulo": "Sistema Sentinela com IA",
  "cliente": "Grupo PARVI", "local": "Manaus/AM", "data": "Julho/2026",
  "slides": [
    {"tipo": "sobre"},                                        # selos padrão (ou "selos":[...])
    {"tipo": "problema", "titulo": "O Desafio", "itens": [{"titulo": "...", "desc": "..."}]},
    {"tipo": "solucao"|"escopo"|"diferenciais", "titulo": "...", "subtitulo": "...", "cards": [{"titulo","desc"}]},
    {"tipo": "passos", "titulo": "Como Funciona", "passos": [{"titulo","desc"}]},
    {"tipo": "kpis", "titulo": "...", "kpis": [{"valor": "+12", "label": "anos"}]},
    {"tipo": "investimento", "titulo": "...", "opcoes": [{"nome","valor","destaque","itens":[...]}], "observacao": "..."},
    {"tipo": "imagem", "titulo": "...", "legenda": "...", "imagem_path": "/caminho.png"},
    {"tipo": "secao", "titulo": "..."},
    {"tipo": "contato", "cta": "Vamos proteger seu pátio?"}   # auto com CEO
  ]
}
A CAPA e o CONTATO são gerados automaticamente (capa do cabeçalho; contato do rodapé/CEO), mas
podem ser sobrescritos por slides explícitos {"tipo":"capa"|"contato"}.
"""
from __future__ import annotations

import io
import os
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from modules.crm.services.pdf_branding import EMPRESA, logo_path

# ── Paleta (mesma dos PDFs) ──────────────────────────────────────────────
LARANJA = RGBColor(0xF9, 0x73, 0x16)
AZUL_ESCURO = RGBColor(0x1E, 0x3A, 0x5F)
AZUL_MEDIO = RGBColor(0x2D, 0x5F, 0x8B)
CINZA = RGBColor(0x1F, 0x29, 0x37)
CINZA_CLARO = RGBColor(0x6B, 0x72, 0x80)
BRANCO = RGBColor(0xFF, 0xFF, 0xFF)
FUNDO = RGBColor(0xF8, 0xFA, 0xFC)
CARD_BORDA = RGBColor(0xE5, 0xE7, 0xEB)
FONTE = "Calibri"

W = Inches(13.333)
H = Inches(7.5)


def _solid(shape, rgb):
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb
    shape.line.fill.background()


def _txt(slide, left, top, width, height, text, size, color, *, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, font=FONTE, italic=False, leading=None):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    lines = text.split("\n")
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if leading:
            p.line_spacing = leading
        r = p.add_run()
        r.text = ln
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.name = font
        r.font.color.rgb = color
    return box


def _rodape(slide, idx: int):
    """Rodapé de marca fixo + numeração (todos os slides internos)."""
    barra = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, H - Inches(0.42), W, Inches(0.42))
    _solid(barra, AZUL_ESCURO)
    marca = f"CONECTA MAIS  ·  CNPJ {EMPRESA['cnpj']}  ·  {EMPRESA['fone']}  ·  {EMPRESA['site']}"
    _txt(slide, Inches(0.5), H - Inches(0.40), Inches(11), Inches(0.35), marca, 9.5, BRANCO,
         anchor=MSO_ANCHOR.MIDDLE)
    _txt(slide, W - Inches(1.2), H - Inches(0.40), Inches(0.9), Inches(0.35), str(idx), 9.5,
         LARANJA, bold=True, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)


def _titulo_slide(slide, titulo: str, subtitulo: str | None = None):
    """Faixa de título padrão: barra lateral laranja + título azul + subtítulo."""
    barra = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(0.5), Inches(0.14), Inches(0.95))
    _solid(barra, LARANJA)
    _txt(slide, Inches(0.8), Inches(0.42), Inches(11.7), Inches(0.7), titulo, 30, AZUL_ESCURO, bold=True)
    if subtitulo:
        _txt(slide, Inches(0.82), Inches(1.12), Inches(11.7), Inches(0.5), subtitulo, 15, LARANJA)


def _fundo(slide, rgb=BRANCO):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    _solid(bg, rgb)
    slide.shapes._spTree.remove(bg._element)
    slide.shapes._spTree.insert(2, bg._element)


def _blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _melhor_logo() -> str | None:
    """Logo de melhor qualidade disponível (a marca é feita p/ fundo claro)."""
    for cand in (
        "/app/uploads/assets/logo-conecta-mais.png",
        "/app/assets/logo.png",
        logo_path("cover"),
        logo_path("header"),
        "/app/uploads/assets/conecta-mais/sublogo-sem-fundo.2.png",
    ):
        if cand and os.path.exists(cand):
            return cand
    return None


def _logo_badge(slide, left, top, logo_h_in: float = 1.05):
    """Logo dentro de um BADGE branco arredondado — para aparecer nítida sobre fundo escuro."""
    cand = _melhor_logo()
    if not cand:
        return
    try:
        from PIL import Image

        with Image.open(cand) as im:
            ratio = im.width / max(im.height, 1)
    except Exception:
        ratio = 1.7
    logo_h = Inches(logo_h_in)
    logo_w = Inches(logo_h_in * ratio)
    pad = Inches(0.26)
    badge_w = logo_w + pad * 2
    badge_h = logo_h + pad * 2
    badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, badge_w, badge_h)
    _solid(badge, BRANCO)
    badge.shadow.inherit = False
    try:
        badge.adjustments[0] = 0.12
    except Exception:
        pass
    try:
        slide.shapes.add_picture(cand, left + pad, top + pad, height=logo_h)
    except Exception:
        pass
    return badge_h


# ── Slides ───────────────────────────────────────────────────────────────
def _slide_capa(prs, dados):
    s = _blank(prs)
    _fundo(s, AZUL_ESCURO)
    # faixa laranja lateral
    faixa = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.35), H)
    _solid(faixa, LARANJA)
    # detalhe geométrico sutil (canto inferior direito)
    deco = s.shapes.add_shape(MSO_SHAPE.OVAL, W - Inches(2.6), H - Inches(2.6), Inches(4.0), Inches(4.0))
    _solid(deco, RGBColor(0x27, 0x46, 0x6E))
    s.shapes._spTree.remove(deco._element)
    s.shapes._spTree.insert(3, deco._element)
    # logo em badge branco (nítida sobre o fundo escuro)
    _logo_badge(s, Inches(0.9), Inches(0.75), logo_h_in=1.05)
    _txt(s, Inches(0.9), Inches(2.5), Inches(11.5), Inches(1.8), dados.get("titulo", ""), 46,
         BRANCO, bold=True, leading=1.05)
    if dados.get("subtitulo"):
        _txt(s, Inches(0.92), Inches(4.2), Inches(11), Inches(0.8), dados["subtitulo"], 22, LARANJA)
    # bloco cliente/local/data
    linha = []
    if dados.get("cliente"):
        linha.append(f"Preparado para: {dados['cliente']}")
    if dados.get("local"):
        linha.append(dados["local"])
    if dados.get("data"):
        linha.append(dados["data"])
    if linha:
        _txt(s, Inches(0.92), Inches(5.3), Inches(11), Inches(1.0), "\n".join(linha), 15,
             RGBColor(0xCB, 0xD5, 0xE1), leading=1.3)
    _txt(s, Inches(0.9), H - Inches(0.7), Inches(11), Inches(0.4),
         f"{EMPRESA['fone']}  ·  {EMPRESA['site']}", 12, LARANJA, bold=True)
    return s


def _cards_grid(s, itens, *, numerado=True, top=Inches(1.9)):
    """Grade de cards (até 6) título+descrição, com número laranja."""
    n = len(itens)
    cols = 2 if n <= 4 else 3
    rows = (n + cols - 1) // cols
    gap = Inches(0.3)
    cw = (W - Inches(1.0) - gap * (cols - 1)) / cols
    ch = min(Inches(1.7), (H - top - Inches(0.7) - gap * (rows - 1)) / rows)
    for i, it in enumerate(itens):
        r, c = divmod(i, cols)
        left = Inches(0.5) + c * (cw + gap)
        tp = top + r * (ch + gap)
        card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, tp, cw, ch)
        card.fill.solid()
        card.fill.fore_color.rgb = FUNDO
        card.line.color.rgb = CARD_BORDA
        card.line.width = Pt(1)
        card.shadow.inherit = False
        tit = it.get("titulo", "")
        prefixo = f"{i + 1}.  " if numerado else ""
        tb = s.shapes.add_textbox(left + Inches(0.25), tp + Inches(0.18), cw - Inches(0.5), ch - Inches(0.36))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        r1 = p.add_run()
        r1.text = prefixo
        r1.font.size = Pt(15)
        r1.font.bold = True
        r1.font.name = FONTE
        r1.font.color.rgb = LARANJA
        r2 = p.add_run()
        r2.text = tit
        r2.font.size = Pt(14.5)
        r2.font.bold = True
        r2.font.name = FONTE
        r2.font.color.rgb = AZUL_ESCURO
        if it.get("desc"):
            pd = tf.add_paragraph()
            pd.space_before = Pt(4)
            rd = pd.add_run()
            rd.text = it["desc"]
            rd.font.size = Pt(11.5)
            rd.font.name = FONTE
            rd.font.color.rgb = CINZA


def _slide_cards(prs, sl, idx, *, numerado=False):
    s = _blank(prs)
    _fundo(s)
    _titulo_slide(s, sl.get("titulo", ""), sl.get("subtitulo"))
    itens = sl.get("cards") or sl.get("itens") or []
    _cards_grid(s, itens, numerado=numerado)
    if sl.get("frase"):
        _txt(s, Inches(0.5), H - Inches(0.95), Inches(12.3), Inches(0.5), sl["frase"], 13,
             AZUL_MEDIO, italic=True, align=PP_ALIGN.CENTER)
    _rodape(s, idx)
    return s


def _slide_passos(prs, sl, idx):
    s = _blank(prs)
    _fundo(s)
    _titulo_slide(s, sl.get("titulo", ""), sl.get("subtitulo"))
    passos = sl.get("passos") or []
    n = len(passos)
    top = Inches(2.1)
    ch = min(Inches(0.95), (H - top - Inches(0.8)) / max(n, 1))
    for i, p in enumerate(passos):
        tp = top + i * (ch + Inches(0.12))
        # círculo com número
        circ = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.7), tp, Inches(0.62), Inches(0.62))
        _solid(circ, LARANJA)
        ctf = circ.text_frame
        ctf.word_wrap = False
        cp = ctf.paragraphs[0]
        cp.alignment = PP_ALIGN.CENTER
        cr = cp.add_run()
        cr.text = str(i + 1)
        cr.font.size = Pt(20)
        cr.font.bold = True
        cr.font.name = FONTE
        cr.font.color.rgb = BRANCO
        tb = s.shapes.add_textbox(Inches(1.6), tp - Inches(0.05), Inches(11.0), ch + Inches(0.1))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        pp = tf.paragraphs[0]
        r1 = pp.add_run()
        r1.text = p.get("titulo", "")
        r1.font.size = Pt(15)
        r1.font.bold = True
        r1.font.name = FONTE
        r1.font.color.rgb = AZUL_ESCURO
        if p.get("desc"):
            r2 = pp.add_run()
            r2.text = "   " + p["desc"]
            r2.font.size = Pt(12.5)
            r2.font.name = FONTE
            r2.font.color.rgb = CINZA
    _rodape(s, idx)
    return s


def _slide_kpis(prs, sl, idx):
    s = _blank(prs)
    _fundo(s)
    _titulo_slide(s, sl.get("titulo", "") or "Números que nos definem", sl.get("subtitulo"))
    kpis = sl.get("kpis") or []
    n = len(kpis)
    gap = Inches(0.3)
    cw = (W - Inches(1.0) - gap * (n - 1)) / max(n, 1)
    top = Inches(2.6)
    for i, k in enumerate(kpis):
        left = Inches(0.5) + i * (cw + gap)
        card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, cw, Inches(2.0))
        _solid(card, AZUL_ESCURO)
        _txt(s, left, top + Inches(0.35), cw, Inches(0.9), str(k.get("valor", "")), 40, LARANJA,
             bold=True, align=PP_ALIGN.CENTER)
        _txt(s, left, top + Inches(1.3), cw, Inches(0.6), str(k.get("label", "")), 14, BRANCO,
             align=PP_ALIGN.CENTER)
    _rodape(s, idx)
    return s


def _slide_investimento(prs, sl, idx):
    s = _blank(prs)
    _fundo(s)
    _titulo_slide(s, sl.get("titulo", "") or "Investimento", sl.get("subtitulo"))
    opcoes = sl.get("opcoes") or []
    n = len(opcoes)
    gap = Inches(0.4)
    cw = min(Inches(4.2), (W - Inches(1.0) - gap * (n - 1)) / max(n, 1))
    total_w = cw * n + gap * (n - 1)
    x0 = (W - total_w) / 2
    top = Inches(2.0)
    for i, o in enumerate(opcoes):
        left = x0 + i * (cw + gap)
        destaque = o.get("destaque")
        card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, cw, Inches(3.9))
        card.fill.solid()
        card.fill.fore_color.rgb = AZUL_ESCURO if destaque else FUNDO
        card.line.color.rgb = LARANJA
        card.line.width = Pt(2.5 if destaque else 1)
        card.shadow.inherit = False
        cor_tit = BRANCO if destaque else AZUL_ESCURO
        cor_txt = RGBColor(0xE5, 0xE7, 0xEB) if destaque else CINZA
        if destaque:
            selo = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left + cw / 2 - Inches(0.9),
                                      top - Inches(0.22), Inches(1.8), Inches(0.44))
            _solid(selo, LARANJA)
            _txt(s, left + cw / 2 - Inches(0.9), top - Inches(0.18), Inches(1.8), Inches(0.36),
                 "MELHOR VALOR", 11, BRANCO, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        _txt(s, left, top + Inches(0.35), cw, Inches(0.5), o.get("nome", ""), 17, cor_tit, bold=True,
             align=PP_ALIGN.CENTER)
        _txt(s, left, top + Inches(1.0), cw, Inches(0.8), o.get("valor", ""), 30, LARANJA, bold=True,
             align=PP_ALIGN.CENTER)
        itens = o.get("itens") or []
        if itens:
            _txt(s, left + Inches(0.3), top + Inches(2.0), cw - Inches(0.6), Inches(1.7),
                 "\n".join(f"• {x}" for x in itens), 11.5, cor_txt, leading=1.25)
    if sl.get("observacao"):
        _txt(s, Inches(0.5), H - Inches(0.95), Inches(12.3), Inches(0.5), sl["observacao"], 12,
             AZUL_MEDIO, italic=True, align=PP_ALIGN.CENTER)
    _rodape(s, idx)
    return s


def _slide_imagem(prs, sl, idx):
    s = _blank(prs)
    _fundo(s)
    _titulo_slide(s, sl.get("titulo", ""), sl.get("subtitulo"))
    ip = sl.get("imagem_path")
    if ip and os.path.exists(ip):
        try:
            s.shapes.add_picture(ip, Inches(1.5), Inches(1.9), width=Inches(10.3))
        except Exception:
            ip = None
    if not (ip and os.path.exists(ip)):
        ph = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(1.9), Inches(10.3), Inches(4.2))
        ph.fill.solid()
        ph.fill.fore_color.rgb = FUNDO
        ph.line.color.rgb = CARD_BORDA
        _txt(s, Inches(1.5), Inches(3.7), Inches(10.3), Inches(0.6), "[ imagem / render do projeto ]",
             14, CINZA_CLARO, align=PP_ALIGN.CENTER)
    if sl.get("legenda"):
        _txt(s, Inches(0.5), H - Inches(0.95), Inches(12.3), Inches(0.5), sl["legenda"], 13,
             AZUL_MEDIO, italic=True, align=PP_ALIGN.CENTER)
    _rodape(s, idx)
    return s


def _slide_secao(prs, sl, idx):
    s = _blank(prs)
    _fundo(s, AZUL_ESCURO)
    faixa = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, H / 2 - Inches(0.06), W, Inches(0.12))
    _solid(faixa, LARANJA)
    _txt(s, Inches(1), H / 2 - Inches(1.1), Inches(11.3), Inches(1.0), sl.get("titulo", ""), 40,
         BRANCO, bold=True, align=PP_ALIGN.CENTER)
    if sl.get("subtitulo"):
        _txt(s, Inches(1), H / 2 + Inches(0.2), Inches(11.3), Inches(0.6), sl["subtitulo"], 18,
             LARANJA, align=PP_ALIGN.CENTER)
    return s


def _slide_sobre(prs, sl, idx):
    selos = sl.get("selos") or [
        {"valor": "+12", "label": "anos de mercado"},
        {"valor": "100%", "label": "equipe própria"},
        {"valor": "24/7", "label": "central de monitoramento"},
        {"valor": "Sentinela", "label": "tecnologia de IA própria"},
    ]
    sl2 = dict(sl)
    sl2.setdefault("titulo", "Quem Somos")
    sl2.setdefault("subtitulo", sl.get("subtitulo") or "Segurança e tecnologia com presença humana onde importa")
    s = _blank(prs)
    _fundo(s)
    _titulo_slide(s, sl2["titulo"], sl2["subtitulo"])
    if sl.get("texto"):
        _txt(s, Inches(0.82), Inches(1.8), Inches(11.7), Inches(1.1), sl["texto"], 14, CINZA, leading=1.3)
    # selos como kpis
    n = len(selos)
    gap = Inches(0.3)
    cw = (W - Inches(1.0) - gap * (n - 1)) / max(n, 1)
    top = Inches(3.3)
    for i, k in enumerate(selos):
        left = Inches(0.5) + i * (cw + gap)
        card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, cw, Inches(2.0))
        _solid(card, AZUL_ESCURO)
        _txt(s, left, top + Inches(0.4), cw, Inches(0.9), str(k.get("valor", "")), 34, LARANJA,
             bold=True, align=PP_ALIGN.CENTER)
        _txt(s, left, top + Inches(1.35), cw, Inches(0.6), str(k.get("label", "")), 13, BRANCO,
             align=PP_ALIGN.CENTER)
    _rodape(s, idx)
    return s


def _slide_contato(prs, sl, idx):
    s = _blank(prs)
    _fundo(s, AZUL_ESCURO)
    faixa = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.35), H)
    _solid(faixa, LARANJA)
    _logo_badge(s, W - Inches(3.3), Inches(0.7), logo_h_in=0.9)
    _txt(s, Inches(0.9), Inches(1.9), Inches(11.5), Inches(1.4),
         sl.get("cta") or "Vamos proteger o que é seu?", 40, BRANCO, bold=True, leading=1.05)
    dados = [
        f"{EMPRESA['fone']}",
        f"{EMPRESA['email']}",
        f"{EMPRESA['site']}",
        f"{EMPRESA['instagram']}",
    ]
    _txt(s, Inches(0.92), Inches(3.2), Inches(8), Inches(1.8), "\n".join(dados), 18,
         RGBColor(0xE5, 0xE7, 0xEB), leading=1.5)
    # assinatura CEO
    _txt(s, Inches(0.92), Inches(5.6), Inches(8), Inches(0.5), EMPRESA["ceo"], 18, LARANJA, bold=True)
    _txt(s, Inches(0.92), Inches(6.05), Inches(8), Inches(0.4),
         f"{EMPRESA['ceo_cargo']} · {EMPRESA['razao']}", 12, RGBColor(0xCB, 0xD5, 0xE1))
    return s


_RENDER = {
    "capa": lambda prs, sl, idx: _slide_capa(prs, sl),
    "sobre": _slide_sobre,
    "problema": lambda prs, sl, idx: _slide_cards(prs, sl, idx, numerado=True),
    "solucao": _slide_cards,
    "escopo": _slide_cards,
    "diferenciais": _slide_cards,
    "cards": _slide_cards,
    "passos": _slide_passos,
    "kpis": _slide_kpis,
    "investimento": _slide_investimento,
    "imagem": _slide_imagem,
    "secao": _slide_secao,
    "contato": _slide_contato,
}


def build_pptx(dados: dict[str, Any]) -> bytes:
    """Gera a apresentação .pptx no padrão Conecta PRO e retorna os bytes."""
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H

    # capa sempre primeiro (do cabeçalho), salvo se o usuário já colocou uma explícita
    slides = list(dados.get("slides") or [])
    if not slides or slides[0].get("tipo") != "capa":
        _slide_capa(prs, dados)
    idx = 1
    tem_contato = any(sl.get("tipo") == "contato" for sl in slides)
    for sl in slides:
        tipo = sl.get("tipo", "cards")
        render = _RENDER.get(tipo, _slide_cards)
        if tipo == "capa":
            render(prs, dados, idx)
        else:
            render(prs, sl, idx)
            idx += 1
    if not tem_contato:
        _slide_contato(prs, {}, idx)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def pptx_to_pdf(pptx_bytes: bytes) -> bytes:
    """Converte .pptx → .pdf com fidelidade via LibreOffice headless (soffice)."""
    import glob
    import subprocess
    import tempfile

    soffice = None
    for cand in ("soffice", "libreoffice", "/usr/bin/soffice", "/usr/bin/libreoffice"):
        try:
            if subprocess.run([cand, "--version"], capture_output=True, timeout=20).returncode == 0:
                soffice = cand
                break
        except Exception:
            continue
    if not soffice:
        raise RuntimeError("LibreOffice (soffice) não encontrado no container")

    with tempfile.TemporaryDirectory() as td:
        pin = os.path.join(td, "apresentacao.pptx")
        with open(pin, "wb") as fh:
            fh.write(pptx_bytes)
        env = dict(os.environ)
        env["HOME"] = td  # soffice precisa de HOME gravável para o profile
        subprocess.run(
            [soffice, "--headless", "--nologo", "--nofirststartwizard",
             "--convert-to", "pdf", "--outdir", td, pin],
            capture_output=True, timeout=180, env=env, check=True,
        )
        pdfs = glob.glob(os.path.join(td, "*.pdf"))
        if not pdfs:
            raise RuntimeError("Conversão pptx→pdf não gerou arquivo")
        with open(pdfs[0], "rb") as fh:
            return fh.read()


def build(dados: dict[str, Any], *, pdf: bool = True) -> dict[str, bytes]:
    """Gera a apresentação. Retorna {'pptx': bytes, 'pdf': bytes?}."""
    out = {"pptx": build_pptx(dados)}
    if pdf:
        try:
            out["pdf"] = pptx_to_pdf(out["pptx"])
        except Exception as e:  # noqa: BLE001 — PDF é best-effort; pptx sempre sai
            out["pdf_error"] = str(e)[:200]
    return out
