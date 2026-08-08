"""Gera a planilha de conferência do PONTO — quem bate 2× e quem bate 4×.

Existe porque a regra não é derivável com segurança do que está no banco: o campo
`employees.recebe_intrajornada` tem erro (4 pessoas marcadas true batem 4× na prática) e a
lista de condomínios que eu tinha usado antes era proxy, não causa. O Jordan é a fonte da
verdade organizacional — esta planilha leva o que o sistema sabe para ele preencher o que
está certo, e volta como carga.

Só LÊ. Não altera cadastro. As colunas que o Jordan preenche vêm vazias e destacadas.

Colunas de evidência (para ele decidir com dado, não de memória):
  • `recebe_intrajornada` — o que o cadastro diz HOJE
  • `batidas/dia (jul+ago)` — o que a pessoa REALMENTE faz
  • `folha jul` — se a Portte pagou a rubrica 0030/0031 em julho
  • `divergência` — onde essas três fontes discordam entre si

Uso:  docker exec conecta-pro-backend python3 /tmp/gerar_planilha_ponto_intrajornada.py
"""
from __future__ import annotations

import os
import sys

sys.path.insert(0, "/app")

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from sqlalchemy import create_engine, text  # noqa: E402

SAIDA = "/tmp/ponto_intrajornada_conferencia.xlsx"

SQL = """
WITH aloc AS (
  SELECT DISTINCT ON (a.employee_id) a.employee_id, c.nome AS cond, coalesce(a.funcao,'—') AS funcao
  FROM employee_alocacoes a JOIN condominios c ON c.id = a.condominio_id
  WHERE a.ativo ORDER BY a.employee_id, a.data_inicio DESC NULLS LAST
), dia AS (
  SELECT employee_id, punch_timestamp::date d, count(*) n
  FROM gp_clock_punches WHERE punch_timestamp >= '2026-07-01' GROUP BY 1, 2
), bat AS (
  SELECT employee_id, round(avg(n), 2) AS media, count(*) AS dias FROM dia GROUP BY 1
), fol AS (
  SELECT DISTINCT employee_id::text AS eid FROM folha_verba_espelho
  WHERE codigo IN ('0030','0031') AND mes = 7 AND ano = 2026
)
SELECT e.nome,
       coalesce(e.cargo, '—')            AS cargo,
       coalesce(al.cond, '(sem alocação)') AS condominio,
       coalesce(al.funcao, '—')          AS funcao_no_posto,
       coalesce(e.escala_padrao, '—')    AS escala,
       coalesce(e.turno_padrao, '—')     AS turno,
       e.recebe_intrajornada             AS flag_cadastro,
       (f.eid IS NOT NULL)               AS folha_julho,
       b.media                           AS batidas_dia,
       coalesce(b.dias, 0)               AS dias_com_batida
FROM employees e
LEFT JOIN aloc al ON al.employee_id = e.id
LEFT JOIN bat  b  ON b.employee_id  = e.id
LEFT JOIN fol  f  ON f.eid = e.id::text
WHERE e.status = 'ativo' AND coalesce(e.is_homologacao, false) = false
ORDER BY coalesce(al.cond, 'zzz'), coalesce(e.cargo,''), e.nome
"""

CABECALHO = [
    ("Colaborador", 34), ("Cargo", 26), ("Condomínio", 20), ("Função no posto", 26),
    ("Escala", 10), ("Turno", 12),
    ("Cadastro diz\nrecebe intrajornada", 18),
    ("Folha jul\npagou intrajornada", 16),
    ("Batidas/dia\n(jul+ago)", 13), ("Dias com\nbatida", 10),
    ("Divergência", 34),
    ("→ RECEBE INTRAJORNADA?\n(preencher: SIM / NAO)", 26),
    ("→ Quantas batidas por turno?\n(preencher: 2 / 4)", 26),
    ("→ Observação", 34),
]

#: colunas que o Jordan preenche
PRIMEIRA_A_PREENCHER = 12


def _divergencia(flag, folha, media) -> str:
    """Onde cadastro, folha e batidas discordam. Vazio = as três contam a mesma história."""
    if media is None:
        return "sem batida no período — não dá para conferir"
    m = float(media)
    bate2 = m < 3.0                      # 2 batidas/dia (com faltas, fica ~1,5–2,5)
    partes = []
    if bool(flag) != bate2:
        partes.append(f"cadastro diz {'SIM' if flag else 'NÃO'} mas bate {m:.1f}×/dia")
    if bool(folha) != bate2:
        partes.append(f"folha {'pagou' if folha else 'não pagou'} mas bate {m:.1f}×/dia")
    if bool(flag) != bool(folha):
        partes.append("cadastro e folha discordam")
    return " · ".join(dict.fromkeys(partes))


def main() -> None:
    db = create_engine(os.environ["DATABASE_URL"].replace("+asyncpg", "")).connect()
    rows = db.execute(text(SQL)).mappings().all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Ponto — intrajornada"

    azul = PatternFill("solid", fgColor="16277D")
    laranja = PatternFill("solid", fgColor="F26522")
    amarelo = PatternFill("solid", fgColor="FEF3C7")
    branco = Font(color="FFFFFF", bold=True, size=10)

    for i, (titulo, larg) in enumerate(CABECALHO, start=1):
        c = ws.cell(row=1, column=i, value=titulo)
        c.fill = laranja if i >= PRIMEIRA_A_PREENCHER else azul
        c.font = branco
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[get_column_letter(i)].width = larg
    ws.row_dimensions[1].height = 34
    ws.freeze_panes = "A2"

    n_div = 0
    for r in rows:
        div = _divergencia(r["flag_cadastro"], r["folha_julho"], r["batidas_dia"])
        if div:
            n_div += 1
        ws.append([
            r["nome"], r["cargo"], r["condominio"], r["funcao_no_posto"],
            r["escala"], r["turno"],
            "SIM" if r["flag_cadastro"] else "NÃO",
            "sim" if r["folha_julho"] else "não",
            float(r["batidas_dia"]) if r["batidas_dia"] is not None else None,
            r["dias_com_batida"],
            div, None, None, None,
        ])
        if div:                                   # linha com divergência fica marcada
            for col in range(1, len(CABECALHO) + 1):
                ws.cell(row=ws.max_row, column=col).fill = amarelo

    # legenda
    ws.append([])
    ws.append(["LEGENDA — as 3 últimas colunas (laranja) são para você preencher."])
    ws.append(["Linha amarela = cadastro, folha e batidas contam histórias diferentes; "
               "é onde vale mais a sua conferência."])
    ws.append(["Batidas/dia perto de 2 = não almoça (recebe intrajornada). "
               "Perto de 4 = almoça uma hora."])
    ws.append(["Média cai abaixo do esperado quando falta batida — por isso o número vem com "
               "os dias, para você ver o peso da amostra."])
    for i in range(4):
        ws.cell(row=ws.max_row - i, column=1).font = Font(italic=True, size=9)

    wb.save(SAIDA)
    print(f"  planilha: {SAIDA}")
    print(f"  {len(rows)} colaboradores ativos · {n_div} com divergência entre as 3 fontes")


if __name__ == "__main__":
    main()
